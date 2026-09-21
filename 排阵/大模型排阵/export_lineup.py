#!/usr/bin/env python3
"""
大模型排阵 - 校验 + Excel 导出
读取 lineup JSON（可手工修改），执行硬校验/目标检查/BT 强度报告，
导出与 2026-09-07-羽毛球排阵.xlsx 相同样式的 Excel。

用法:
    python export_lineup.py [lineup.json 路径]
    （默认读取本目录下 2026-09-14-lineup.json）
"""

import json
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

BASE_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(BASE_DIR.parent))
from excel_exporter import (  # noqa: E402
    INTERNAL_MALE_PLAYERS, INTERNAL_FEMALE_PLAYERS,
    GUEST_MALE_PLAYERS, GUEST_FEMALE_PLAYERS,
    MALE_PLAYERS, FEMALE_PLAYERS,
)

INTERNAL_PLAYERS = set(INTERNAL_MALE_PLAYERS + INTERNAL_FEMALE_PLAYERS)
GUEST_PLAYERS = set(GUEST_MALE_PLAYERS + GUEST_FEMALE_PLAYERS)
ALL_KNOWN = INTERNAL_PLAYERS | GUEST_PLAYERS
MALES = set(MALE_PLAYERS)
FEMALES = set(FEMALE_PLAYERS)

# 周实例配置（名单/锚点/训练目标 随每周接龙变化，改 instances/<date>.json）
sys.path.insert(0, str(BASE_DIR))
from weekly import load as _load_week  # noqa: E402
WEEK = _load_week()

# BT θ 值（实力乘子），来源 scores/20260831/BT排名总表.xlsx
# 运行时若 xlsx 存在则优先从 xlsx 读取
BT_XLSX = BASE_DIR.parent.parent / "scores" / "20260831" / "BT排名总表.xlsx"
THETA_FALLBACK = {
    "刘海锐": 6.4914, "黄冬青": 3.9086, "林锋": 3.6337, "苏大哲": 2.2303,
    "程建兴": 2.1929, "陈财贵": 2.1839, "张燕红": 2.0617, "林琪琛": 1.8123,
    "刘继宇": 1.4550, "项小英": 1.1782, "范智强": 1.2773, "罗蒙": 1.0456,
    "董广博": 1.0209, "罗琴荩": 1.0083, "严勇文": 0.8907, "陈小洪": 0.8257,
    "唐英武": 0.7730, "王小波": 0.7494, "陈顺星": 0.7343, "王苏丹": 0.7246,
    "卢志辉": 0.6663, "崔倩男": 0.6539, "田茜": 0.6415, "徐越": 0.6033,
    "张欣欣": 0.4835, "李佳琳": 0.5200, "李祺祺": 0.3526, "滕菲": 0.3047,
    "林小连": 0.3033, "谢卓珊": 0.1745,
}


def load_theta():
    theta = dict(THETA_FALLBACK)
    try:
        wb = openpyxl.load_workbook(BT_XLSX, read_only=True, data_only=True)
        for sheet in ("BT排名（出场≥30局）", "全部选手（含样本不足）"):
            if sheet not in wb.sheetnames:
                continue
            ws = wb[sheet]
            for row in ws.iter_rows(values_only=True):
                if not row or len(row) < 4:
                    continue
                name = row[2] if sheet == "BT排名（出场≥30局）" else row[2]
                val = row[3]
                if isinstance(name, str) and name in ALL_KNOWN and isinstance(val, (int, float)):
                    theta[name] = float(val)
        wb.close()
    except Exception as e:
        print(f"  （BT 排名 xlsx 读取失败，使用内置 θ 值：{e}）")
    return theta


def theta(player, theta_map):
    return theta_map.get(player, 1.0)


def win_prob(a, b):
    if a + b <= 0:
        return 0.5
    return a / (a + b)


def match_players(m):
    return set(m["a"]) | set(m["b"])


def validate(lineup, theta_map):
    errors, warnings = [], []
    rounds = lineup["rounds"]
    courts = lineup["courts"]
    caps = lineup.get("caps", WEEK.caps)

    player_total = defaultdict(int)
    player_energy = defaultdict(int)
    player_type = defaultdict(lambda: defaultdict(int))
    pair_count = defaultdict(int)
    type_count = defaultdict(int)

    for rd in rounds:
        rn = rd["round"]
        seen_players = set()
        seen_courts = set()
        for m in rd["matches"]:
            court, mtype = m["court"], m["type"]
            a, b = m["a"], m["b"]
            players = match_players(m)

            if court in seen_courts:
                errors.append(f"第{rn}轮：场地{court}号重复使用")
            seen_courts.add(court)
            if court not in courts:
                errors.append(f"第{rn}轮：场地{court}号不在配置 {courts} 中")

            # 类型与性别合法性
            if mtype == "女单":
                if len(a) != 1 or len(b) != 1 or not set(a) <= FEMALES or not set(b) <= FEMALES:
                    errors.append(f"第{rn}轮{court}号：女单必须是 1名女性 vs 1名女性 -> {a} vs {b}")
            elif mtype == "男单":
                if len(a) != 1 or len(b) != 1 or not set(a) <= MALES or not set(b) <= MALES:
                    errors.append(f"第{rn}轮{court}号：男单必须是 1名男性 vs 1名男性 -> {a} vs {b}")
            elif mtype == "男双":
                if len(a) != 2 or len(b) != 2 or not (set(a) | set(b)) <= MALES:
                    errors.append(f"第{rn}轮{court}号：男双必须是 2男 vs 2男 -> {a} vs {b}")
            elif mtype == "女双":
                if len(a) != 2 or len(b) != 2 or not (set(a) | set(b)) <= FEMALES:
                    errors.append(f"第{rn}轮{court}号：女双必须是 2女 vs 2女 -> {a} vs {b}")
            elif mtype == "混双":
                for side in (a, b):
                    if len(side) != 2 or not (set(side) & MALES) or not (set(side) & FEMALES):
                        errors.append(f"第{rn}轮{court}号：混双每边必须 1男+1女 -> {a} vs {b}")
                        break
            elif mtype == "女双vs混双":
                if len(a) != 2 or not set(a) <= FEMALES:
                    errors.append(f"第{rn}轮{court}号：女双vs混双 女双侧必须 2女 -> {a} vs {b}")
                elif len(b) != 2 or not (set(b) & MALES) or not (set(b) & FEMALES):
                    errors.append(f"第{rn}轮{court}号：女双vs混双 混双侧必须 1男+1女 -> {a} vs {b}")
            else:
                errors.append(f"第{rn}轮{court}号：未知类型 {mtype}")

            # 名单内
            unknown = players - ALL_KNOWN
            if unknown:
                errors.append(f"第{rn}轮{court}号：球员不在名单内 {unknown}")

            # 同一轮重复
            dup = seen_players & players
            if dup:
                errors.append(f"第{rn}轮：球员同轮重复出现 {dup}")
            seen_players |= players

            # 外援规则：整场不能无内部员工（外援不内讧）
            if not (players & INTERNAL_PLAYERS):
                errors.append(f"第{rn}轮{court}号：整场无内部员工（外援内讧） {sorted(players)}")
            # 无外援PK外援：每场至多 1 队含外援（外援+外援同队允许）
            ga, gb = set(a) & GUEST_PLAYERS, set(b) & GUEST_PLAYERS
            if ga and gb:
                errors.append(
                    f"第{rn}轮{court}号：同场外援PK外援 {sorted(ga | gb)}")
            # 双打：同一队 2 人都是外援时，对面必须有内部（players & INTERNAL 已保证，但对面需检查）
            if mtype in ("男双", "女双", "混双"):
                for i, side in enumerate((a, b)):
                    if set(side) <= GUEST_PLAYERS:
                        other = b if i == 0 else a
                        if not (set(other) & INTERNAL_PLAYERS):
                            errors.append(
                                f"第{rn}轮{court}号：外援队 {sorted(side)} 对面无内部员工 {sorted(other)}")

            # 统计（能量消耗=单打×2+双打×1，用户口径）
            inc = 2 if mtype in ("男单", "女单") else 1
            for p in players:
                player_total[p] += 1
                player_energy[p] += inc
                player_type[p][mtype] += 1
            type_count[mtype] += 1
            for side in (a, b):
                if len(side) == 2:
                    pair_count[tuple(sorted(side))] += 1

            # 场次上限
            for p in players:
                limit = caps["guest"] if p in GUEST_PLAYERS else caps["internal"]
                if player_total[p] > limit:
                    errors.append(f"第{rn}轮{court}号：{p} 总场次 {player_total[p]} 超过上限 {limit}")

        # 锚点对完整性：两人同轮打男双必须同队（不拆）
        # ⚠️ 重组轮豁免：该轮锚点本就拆开重组（试搭，如 陈财贵+刘继宇 vs 苏大哲+严勇文）
        if rn not in WEEK.reshuffle_set:
            def md_side_of(name):
                for mi, mm in enumerate(rd["matches"]):
                    if mm["type"] != "男双":
                        continue
                    for si, side in enumerate((mm["a"], mm["b"])):
                        if name in side:
                            return (mi, si)
                return None
            for p1, p2 in WEEK.anchor_pairs:
                loc1, loc2 = md_side_of(p1), md_side_of(p2)
                if loc1 and loc2 and loc1 != loc2:
                    errors.append(f"第{rn}轮：{p1}/{p2} 被拆开（男双不同队）")

    # 同队上限（换搭档，硬）：指定组合同队次数不得超过 max
    def _pk(a, b):
        return tuple(sorted([a, b]))
    for item in WEEK.max_pairs:
        p1, p2 = item["pair"]
        n = pair_count.get(_pk(p1, p2), 0)
        if n > item["max"]:
            errors.append(f"{p1}/{p2} 同队 {n} 次（上限 {item['max']}，换搭档未满足）")
    # 每人场次下限 mins（硬）
    for p, mn in WEEK.mins.items():
        n = player_total.get(p, 0)
        if n < mn:
            errors.append(f"{p} 总场次 {n}（下限 ≥{mn}）")

    # 目标检查（警告级，全部来自周实例 goals）
    def pair_key(p1, p2):
        return tuple(sorted([p1, p2]))
    g = WEEK.goals
    for item in g.get("anchor_pair_min", []):
        p1, p2 = item["pair"]
        n = pair_count.get(pair_key(p1, p2), 0)
        if n < item["min"]:
            warnings.append(f"锚点 {p1}/{p2} 合阵 {n} 场（目标 ≥{item['min']}）")
    for item in g.get("xd_min", []):
        n = player_type[item["player"]].get("混双", 0)
        if n < item["min"]:
            warnings.append(f"{item['player']} 混双 {n} 场（目标 ≥{item['min']}）")
    if g.get("ws_count") is not None and type_count.get("女单", 0) != g["ws_count"]:
        warnings.append(f"女单 {type_count.get('女单', 0)} 场（目标 {g['ws_count']}）")
    if g.get("seed_ms_min") is not None:
        for p in WEEK.seeds:
            n = player_type[p].get("男单", 0)
            if n < g["seed_ms_min"]:
                warnings.append(f"{p} 男单 {n} 场（目标 ≥{g['seed_ms_min']}）")

    # BT 强度报告
    print("\n【BT 强度报告】（双打=θ乘积，单打=θ；>75% 视为偏悬殊）")
    lopsided = []
    for rd in rounds:
        for m in rd["matches"]:
            mtype = m["type"]
            if mtype in ("男单", "女单"):
                sa, sb = theta(m["a"][0], theta_map), theta(m["b"][0], theta_map)
            else:
                sa = prod_theta(m["a"], theta_map)
                sb = prod_theta(m["b"], theta_map)
            p = win_prob(sa, sb)
            fav = m["a"] if p >= 0.5 else m["b"]
            fav_p = max(p, 1 - p)
            name = lambda side: "/".join(side)
            flag = " ⚠️ 偏悬殊" if fav_p > 0.75 else ""
            if fav_p > 0.75:
                lopsided.append((rd["round"], m["court"], mtype, name(m["a"]), name(m["b"]), fav_p))
            print(f"  R{rd['round']} {m['court']}号 {mtype}: "
                  f"{name(m['a'])}({sa:.2f}) vs {name(m['b'])}({sb:.2f}) "
                  f"→ 胜方 {name(fav)} {fav_p:.0%}{flag}")

    # 场次/能量统计（按能量消耗降序；能量=单打×2+双打×1）
    print("\n【球员场次/能量统计】（能量=单打×2+双打×1）")
    for p in sorted(player_total, key=lambda x: (-player_energy[x], -player_total[x], x)):
        t = player_type[p]
        tag = "外" if p in GUEST_PLAYERS else "内"
        detail = " ".join(f"{k}{v}" for k, v in t.items() if v)
        print(f"  {p}({tag}): {player_total[p]}场 / 能量{player_energy[p]}  {detail}")

    type_str = "  ".join(f"{k}{v}" for k, v in sorted(type_count.items()))
    print(f"\n【比赛类型】{type_str}  共{sum(type_count.values())}场")
    return errors, warnings, lopsided


def prod_theta(side, theta_map):
    v = 1.0
    for p in side:
        v *= theta(p, theta_map)
    return v


def export_excel(lineup, output_path, theta_map):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "对阵表"

    font = Font(name="微软雅黑", size=14)
    font_bold = Font(name="微软雅黑", size=14, bold=True)
    center = Alignment(horizontal="center", vertical="center")
    border = Border(left=Side(style="thin"), right=Side(style="thin"),
                    top=Side(style="thin"), bottom=Side(style="thin"))
    header_fill = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid")
    col_a_fill = PatternFill(start_color="EEF2F7", end_color="EEF2F7", fill_type="solid")
    rest_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")

    ws.merge_cells("A1:G1")
    ws["A1"] = lineup["title"]
    ws["A1"].font = font_bold
    ws["A1"].alignment = center

    headers = ["轮次", "场地", "类型", "对阵A", "比分A", "比分B", "对阵B", "胜率参考"]
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=2, column=col, value=h)
        c.font = font_bold
        c.alignment = center
        c.fill = header_fill
        c.border = border

    row = 3
    rounds = lineup["rounds"]
    for rd in rounds:
        for m in rd["matches"]:
            ws.cell(row=row, column=1, value=rd["round"])
            ws.cell(row=row, column=2, value=f"{m['court']}号")
            ws.cell(row=row, column=3, value=m["type"])
            ws.cell(row=row, column=4, value="/".join(m["a"]))
            ws.cell(row=row, column=5, value="")
            ws.cell(row=row, column=6, value="")
            ws.cell(row=row, column=7, value="/".join(m["b"]))
            if m["type"] in ("男单", "女单"):
                sa = theta_map.get(m["a"][0], 1.0)
                sb = theta_map.get(m["b"][0], 1.0)
            else:
                sa = prod_theta(m["a"], theta_map)
                sb = prod_theta(m["b"], theta_map)
            a_pct = int(round(win_prob(sa, sb) * 100))
            ws.cell(row=row, column=8, value=f"{a_pct}:{100 - a_pct}")
            for col in range(1, 9):
                c = ws.cell(row=row, column=col)
                c.font = font
                c.alignment = center
                c.border = border
            ws.cell(row=row, column=1).fill = col_a_fill
            row += 1

    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["D"].width = 17.75
    ws.column_dimensions["E"].width = 8.5
    ws.column_dimensions["G"].width = 17.75
    ws.column_dimensions["H"].width = 10.5
    ws.page_setup.paperSize = 9
    ws.page_setup.orientation = "portrait"

    # 球员局数统计 sheet
    ws2 = wb.create_sheet("球员统计")
    font14 = Font(name="微软雅黑", size=14)
    font14b = Font(name="微软雅黑", size=14, bold=True)
    center2 = Alignment(horizontal="center", vertical="center")
    border2 = Border(left=Side(style="thin"), right=Side(style="thin"),
                     top=Side(style="thin"), bottom=Side(style="thin"))

    games = defaultdict(int)
    for rd in rounds:
        for m in rd["matches"]:
            inc = 2 if m["type"] in ("男单", "女单") else 1
            for p in match_players(m):
                games[p] += inc

    ws2.merge_cells("A1:C1")
    ws2["A1"] = "球员能量消耗统计（能量=单打×2+双打×1）"
    ws2["A1"].font = font14b
    ws2["A1"].alignment = center2
    for col, h in enumerate(["队员", "性别", "能量消耗"], 1):
        c = ws2.cell(row=2, column=col, value=h)
        c.font = font14b
        c.alignment = center2
        c.border = border2

    row2 = 3
    for p in sorted(games, key=lambda x: (-games[x], x)):
        ws2.cell(row=row2, column=1, value=p).font = font14
        ws2.cell(row=row2, column=2, value="男" if p in MALES else "女").font = font14
        ws2.cell(row=row2, column=3, value=games[p]).font = font14
        for col in range(1, 4):
            ws2.cell(row=row2, column=col).alignment = center2
            ws2.cell(row=row2, column=col).border = border2
        row2 += 1

    ws2.column_dimensions["A"].width = 18
    ws2.column_dimensions["B"].width = 8
    ws2.column_dimensions["C"].width = 9

    # 按人 sheet：每人一个 tab，点开即见自己的全部比赛
    n_sheets = export_player_sheets(wb, lineup, theta_map)

    wb.save(output_path)
    print(f"\n✓ Excel 已生成：{output_path}（含 {n_sheets} 个按人 sheet）")


def export_player_sheets(wb, lineup, theta_map):
    """为每名出场球员建一个 sheet：R1-R8 逐轮列出（轮休标注）+ 我方胜率。"""
    seen = set()
    players_in = []
    for rd in lineup["rounds"]:
        for m in rd["matches"]:
            for p in match_players(m):
                if p not in seen:
                    seen.add(p)
                    players_in.append(p)
    # 内部按 BT 排名(θ)降序在前（最关注），外援随后（也按 θ 降序）
    ordered = sorted(seen,
                     key=lambda p: (1 if p in GUEST_PLAYERS else 0,
                                    -theta_map.get(p, 0.0)))

    font = Font(name="微软雅黑", size=14)
    font_bold = Font(name="微软雅黑", size=14, bold=True)
    center = Alignment(horizontal="center", vertical="center")
    border = Border(left=Side(style="thin"), right=Side(style="thin"),
                    top=Side(style="thin"), bottom=Side(style="thin"))
    header_fill = PatternFill(start_color="D9EAD3", end_color="D9EAD3",
                              fill_type="solid")
    col_a_fill = PatternFill(start_color="EEF2F7", end_color="EEF2F7",
                             fill_type="solid")
    rest_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9",
                            fill_type="solid")

    for p in ordered:
        ws = wb.create_sheet(p)
        ws.merge_cells("A1:F1")
        ws["A1"] = f"{p} · {lineup['date']} 个人赛程"
        ws["A1"].font = font_bold
        ws["A1"].alignment = center
        headers = ["轮次", "场地", "类型", "我方", "对手", "我方胜率(参考)"]
        for col, h in enumerate(headers, 1):
            c = ws.cell(row=2, column=col, value=h)
            c.font = font_bold
            c.alignment = center
            c.fill = header_fill
            c.border = border

        row = 3
        n_matches = 0
        n_energy = 0
        for rd in lineup["rounds"]:
            rn = rd["round"]
            match = next((m for m in rd["matches"] if p in match_players(m)), None)
            if match:
                n_matches += 1
                n_energy += 2 if match["type"] in ("男单", "女单") else 1
                on_a = p in match["a"]
                ours = match["a"] if on_a else match["b"]
                theirs = match["b"] if on_a else match["a"]
                if match["type"] in ("男单", "女单"):
                    sa = theta_map.get(match["a"][0], 1.0)
                    sb = theta_map.get(match["b"][0], 1.0)
                else:
                    sa = prod_theta(match["a"], theta_map)
                    sb = prod_theta(match["b"], theta_map)
                a_pct = int(round(win_prob(sa, sb) * 100))
                our_pct = a_pct if on_a else 100 - a_pct
                vals = [rn, f"{match['court']}号", match["type"],
                        "/".join(ours), "/".join(theirs), f"{our_pct}%"]
            else:
                vals = [rn, "", "轮休", "", "", ""]
            for col, v in enumerate(vals, 1):
                c = ws.cell(row=row, column=col, value=v)
                c.font = font
                c.alignment = center
                c.border = border
            if match:
                ws.cell(row=row, column=1).fill = col_a_fill
            else:
                for col in range(1, 7):
                    ws.cell(row=row, column=col).fill = rest_fill
            row += 1

        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        c = ws.cell(row=row, column=1,
                    value=f"共 {n_matches} 场 / 能量消耗 {n_energy}（单打×2+双打×1）　｜　胜率为 BT 模型估算，仅供参考")
        c.font = font
        c.alignment = center
        c.fill = rest_fill

        ws.column_dimensions["A"].width = 6
        ws.column_dimensions["B"].width = 8
        ws.column_dimensions["C"].width = 8
        ws.column_dimensions["D"].width = 18
        ws.column_dimensions["E"].width = 18
        ws.column_dimensions["F"].width = 15
    return len(ordered)


def main():
    default = BASE_DIR / f"{WEEK.date}-lineup.json"
    path = Path(sys.argv[1]) if len(sys.argv) > 1 else default
    with open(path, encoding="utf-8") as f:
        lineup = json.load(f)

    print("=" * 60)
    print(f"大模型排阵校验：{lineup['title']}")
    print("=" * 60)

    theta_map = load_theta()
    errors, warnings, lopsided = validate(lineup, theta_map)

    print("\n【目标达成】")
    def pk(p1, p2):
        return tuple(sorted([p1, p2]))
    target = {f"锚点 {item['pair'][0]}/{item['pair'][1]} 合阵":
              (pk(*item["pair"]), item["min"])
             for item in WEEK.goals.get("anchor_pair_min", [])}
    for label, (key, want) in target.items():
        n = sum(1 for rd in lineup["rounds"] for m in rd["matches"]
                if (len(m["a"]) == 2 and pk(*m["a"]) == key)
                or (len(m["b"]) == 2 and pk(*m["b"]) == key))
        print(f"  {label}：{n} 场（目标 ≥{want}）{'✓' if n >= want else '⚠️'}")
    for item in WEEK.pairs:
        p1, p2 = item["pair"]
        n = sum(1 for rd in lineup["rounds"] for m in rd["matches"]
                if m["type"] == "男双" and (set(item["pair"]) <= set(m["a"])
                                             or set(item["pair"]) <= set(m["b"])))
        print(f"  合练对 {p1}/{p2} 同队：{n} 场（目标 ≥{item['min']}）"
              f"{'✓' if n >= item['min'] else '⚠️'}")
    # 试搭（重组轮）：锚点拆开重组，如 陈财贵+刘继宇
    if WEEK.reshuffle_teams:
        t1 = frozenset(WEEK.reshuffle_teams[0])
        n = sum(1 for rd in lineup["rounds"] for m in rd["matches"]
                if m["type"] == "男双" and (set(m["a"]) == t1 or set(m["b"]) == t1))
        want = len(WEEK.reshuffle_rounds)
        rr = "/".join("R" + str(r) for r in sorted(WEEK.reshuffle_rounds))
        print(f"  试搭 {'/'.join(WEEK.reshuffle_teams[0])}：{n} 场（目标 {want}，重组轮 {rr}）"
              f"{'✓' if n >= want else '⚠️'}")
    # 每人场次下限（mins）+ 同队上限（换搭档）
    for p, mn in WEEK.mins.items():
        n = sum(1 for rd in lineup["rounds"] for m in rd["matches"]
                if p in set(m["a"]) | set(m["b"]))
        print(f"  场次下限 {p}：{n} 场（目标 ≥{mn}）{'✓' if n >= mn else '⚠️'}")
    for item in WEEK.max_pairs:
        p1, p2 = item["pair"]
        fk = frozenset(item["pair"])
        n = sum(1 for rd in lineup["rounds"] for m in rd["matches"]
                if m["type"] == "男双" and (set(m["a"]) == fk or set(m["b"]) == fk))
        print(f"  换搭档 {p1}/{p2} 同队：{n} 次（上限 {item['max']}）"
              f"{'✓' if n <= item['max'] else '⚠️'}")
    # 互搞矩阵：非锚点合练对 对上了哪些锚点队（软目标：尽量两个锚点队都遇上）
    anchor_sets = {frozenset(a) for a in WEEK.anchor_pairs}
    nonanchor_pairs = [it for it in WEEK.pairs if frozenset(it["pair"]) not in anchor_sets]
    for it in nonanchor_pairs:
        pk = frozenset(it["pair"])
        faced = {}
        for rd in lineup["rounds"]:
            for m in rd["matches"]:
                if m["type"] != "男双":
                    continue
                for i, side in enumerate((m["a"], m["b"])):
                    if frozenset(side) == pk:
                        other = (m["b"], m["a"])[i]
                        for ap in WEEK.anchor_pairs:
                            if frozenset(other) == frozenset(ap):
                                aname = "/".join(ap)
                                faced[aname] = faced.get(aname, 0) + 1
        detail = "、".join(f"{k}×{v}" for k, v in faced.items()) or "未对上锚点队"
        ok = len(faced) >= len(WEEK.anchor_pairs)
        print(f"  互搞 {'/'.join(it['pair'])} → {detail}"
              f"{'✓(对上全部锚点队)' if ok else '⚠️(未对上全部锚点队)'}")
    # 种子单打对手尽量不同人（软目标：风格多样，0 重复=全不同）
    seed_opp = {}
    for rd in lineup["rounds"]:
        for m in rd["matches"]:
            if m["type"] == "男单":
                a, b = m["a"][0], m["b"][0]
                if a in WEEK.seeds:
                    seed_opp.setdefault(a, []).append(b)
                elif b in WEEK.seeds:
                    seed_opp.setdefault(b, []).append(a)
    dup = 0
    for opps in seed_opp.values():
        c = {}
        for o in opps:
            c[o] = c.get(o, 0) + 1
        dup += sum(x - 1 for x in c.values() if x > 1)
    print(f"  种子单打对手重复：{dup} 处（0=每个对手都不同，风格多样）"
          f"{'✓' if dup == 0 else ''}")

    for w in warnings:
        print(f"  ⚠️ {w}")
    if not warnings:
        print("  ✓ 全部训练目标达成")

    if errors:
        print(f"\n❌ 硬校验失败（{len(errors)} 处）：")
        for e in errors:
            print(f"  ✗ {e}")
        sys.exit(1)
    print("\n✓ 硬校验全部通过（同轮不重复/场地/类型性别/名单/场次上限/"
          "锚点不拆/无外援PK外援）")
    if lopsided:
        print(f"  （提示：{len(lopsided)} 场强度偏悬殊，可接受：锚点队/种子内战碾压场）")

    out = BASE_DIR / f"{lineup['date']}-羽毛球排阵.xlsx"
    export_excel(lineup, out, theta_map)


if __name__ == "__main__":
    main()
